"""Build tests/fixtures/Rental_Agreement_Sample_04.docx.

Reconstructed regression fixture (SYNTHETIC — all parties/amounts fictional)
exhibiting the symptoms from the observed 10 April 2026 diagnosis:
numbered clauses, BY AND BETWEEN preamble, ownership declaration, an
eleven-month term clause, ordinary rent/deposit/utility clauses, and
signature + witness blocks.
"""
from pathlib import Path

from docx import Document as DocxDocument

FIXTURE = [
    ("title", "RENTAL AGREEMENT"),
    ("p", "This Rental Agreement is made and executed on 10 April 2026 at Bengaluru, Karnataka,"),
    ("p", "BY AND BETWEEN:"),
    ("p", "Mr. Ramesh Iyer, aged 45 years, residing at Flat 402, Shanti Apartments, 4th Cross, Jayanagar, Bengaluru (hereinafter referred to as the \"Owner\" / \"Landlord\")"),
    ("p", "AND"),
    ("p", "Ms. Divya Nair, aged 32 years, residing at No. 17, Rose Garden Layout, HSR Sector 2, Bengaluru (hereinafter referred to as the \"Tenant\")"),
    ("p", "NOW THIS AGREEMENT WITNESSETH AS FOLLOWS:"),
    ("p", "1. The Landlord has agreed to let out the residential Flat No. 402, 4th Floor, Shanti Apartments, Jayanagar, Bengaluru, measuring about 950 sq. ft., together with one covered car parking space, to the Tenant for residential use."),
    ("p", "2. The Landlord represents and declares that he is the absolute owner of the said premises, that the premises are free from all encumbrances, and that he has full authority to grant this tenancy."),
    ("p", "3. The Tenant shall pay a monthly rent of Rs. 18,000 (Rupees Eighteen Thousand only), payable on or before the 5th day of each calendar month by bank transfer."),
    ("p", "4. This agreement shall remain in force for eleven (11) months from the execution date, unless terminated earlier in accordance with this agreement. Renewal may be made by mutual written consent."),
    ("p", "5. The Tenant shall pay a refundable security deposit of Rs. 50,000 (Rupees Fifty Thousand only). The deposit shall be refunded within 30 days of vacating the premises after deducting documented outstanding dues."),
    ("p", "6. The Tenant shall pay for electricity and water consumed in the premises. Property tax and society maintenance charges shall be borne by the Landlord."),
    ("p", "7. The premises shall be used only for residential purposes by the Tenant and her immediate family members."),
    ("p", "8. The Landlord shall carry out all major structural repairs. The Tenant shall maintain interior fixtures in good condition."),
    ("p", "9. The Landlord may enter the premises for inspection with 24 hours prior written notice during daytime."),
    ("p", "10. Either party may terminate this agreement by giving 30 days written notice to the other party."),
    ("p", "11. If rent remains unpaid for more than 7 days after the due date, a late fee of Rs. 500 shall apply for each month of delay."),
    ("p", "12. The Tenant shall not sublet the premises without the prior written consent of the Landlord, which shall not be unreasonably withheld."),
    ("p", "13. Any dispute arising from this agreement shall be referred to arbitration under the Arbitration and Conciliation Act, 1996 at Bengaluru."),
    ("p", "14. This agreement constitutes the entire understanding between the parties and supersedes all prior discussions."),
    ("p", "IN WITNESS WHEREOF, the parties have set their hands on the date first written above."),
    ("p", "Signature of Landlord: ________________      Date: __________"),
    ("p", "Signature of Tenant: ________________      Date: __________"),
    ("p", "WITNESS 1: Name: __________  Address: __________  Signature: __________"),
    ("p", "WITNESS 2: Name: __________  Address: __________  Signature: __________"),
]

out = Path(__file__).resolve().parent / "Rental_Agreement_Sample_04.docx"
doc = DocxDocument()
for kind, text in FIXTURE:
    if kind == "title":
        doc.add_heading(text, level=1)
    else:
        doc.add_paragraph(text)
doc.save(out)
print(f"Wrote {out} ({len(FIXTURE)} blocks)")
