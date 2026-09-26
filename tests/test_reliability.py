"""
Reliability regression tests for the Sample_04 work order.

Covers: preamble exclusion, heading contamination, topic classification
(ownership vs rent, signatures, mutual vs unilateral, penalties),
eleven-month expiry math, neutralising-variant gating, illustrative
litigation disclosure, provenance, unclassified behaviour, compounds,
DOCX tables, phylogeny orientation, provisional-leaf lineage safety, and a
full end-to-end diagnosis of the reconstructed Sample_04 fixture with a
deterministic offline embedding stub (no model required).

The fixture tests/fixtures/Rental_Agreement_Sample_04.docx is SYNTHETIC and
reconstructed to exhibit the reported 10 April 2026 failure symptoms.
"""
from __future__ import annotations

import asyncio
import json
import sys
import uuid
from datetime import date
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from sqlmodel import Session, select

from strain.backend.pipeline.analyse import (
    _normalised_edit_distance,
    assess_asymmetry,
    assess_harshness_vs_root,
    assess_outcome_factor,
    detect_mutation,
    extract_clause_features,
    orient_edge,
)
from strain.backend.pipeline.diagnose import (
    NO_VARIANT_MESSAGE,
    _add_duration,
    _build_handoff_panel,
    _compare_material_terms,
    _diagnose_clause,
    _extract_typed_dates,
    _get_outcomes_for_strain,
    _select_reference_variant,
    diagnose,
)
from strain.backend.pipeline.embed import OfflineEmbeddingProvider
from strain.backend.pipeline.segment import (
    TOPIC_LABELS,
    ClauseRecord,
    classify_topic,
    is_compound,
    normalise,
    read_document,
    segment,
)
from strain.backend.store.store import Clause, Document, Edge, Strain, engine

FIXTURE_DOCX = Path(__file__).resolve().parent / "fixtures" / "Rental_Agreement_Sample_04.docx"

TERM_11 = (
    "This agreement shall remain in force for eleven (11) months from the "
    "execution date, unless terminated earlier in accordance with this "
    "agreement. Renewal may be made by mutual written consent."
)
TERM_3_BAD = (
    "This agreement shall remain in force for THREE (3) MONTHS from 15 March 2011. "
    "Renewal requires seven days written notice before expiry."
)


def sample04_records() -> list[ClauseRecord]:
    assert FIXTURE_DOCX.exists(), "Sample_04 fixture missing"
    return segment(read_document(FIXTURE_DOCX), "doc-sample04")


def operative(recs: list[ClauseRecord]) -> list[ClauseRecord]:
    return [r for r in recs if r.kind == "operative"]


# ─── 1. Preamble exclusion ────────────────────────────────────────────────

class TestPreambleExclusion:
    def test_exactly_one_preamble(self):
        recs = sample04_records()
        pres = [r for r in recs if r.kind == "preamble"]
        assert len(pres) == 1

    def test_preamble_holds_parties_not_scored(self):
        recs = sample04_records()
        pre = next(r for r in recs if r.kind == "preamble")
        assert "Ramesh" in pre.text or "BY AND BETWEEN" in pre.text

    def test_no_clause_text_starts_with_preamble(self):
        for r in operative(sample04_records()):
            assert not r.text.startswith("BY AND BETWEEN")
            assert "hereinafter referred to as" not in r.text


# ─── 2. Heading contamination ──────────────────────────────────────────────

class TestHeadingHygiene:
    def test_no_by_and_between_headings(self):
        for r in sample04_records():
            assert "by and between" not in r.heading.lower(), r.heading

    def test_no_running_doc_title(self):
        heads = [r.heading for r in operative(sample04_records())]
        assert not any(h.strip().lower() == "rental agreement" for h in heads)

    def test_headings_short_or_empty(self):
        for r in operative(sample04_records()):
            assert len(r.heading) <= 60, f"Heading leaked body text: {r.heading!r}"


# ─── 3/4. Topic classification ─────────────────────────────────────────────

class TestTopics:
    def test_ownership_is_not_rent(self):
        recs = sample04_records()
        own = next(r for r in recs if "absolute owner" in r.text)
        assert own.topic == "ownership_authority", own.topic

    def test_all_expected_topics_present(self):
        topics = {r.topic for r in operative(sample04_records())}
        for t in ("property_schedule", "ownership_authority", "rent_payment",
                  "term_renewal", "deposit_refund", "utilities", "permitted_use",
                  "maintenance_repairs", "landlord_access", "termination_notice",
                  "late_penalty", "subletting", "dispute_resolution"):
            assert t in topics, f"Missing topic {t}; got {sorted(topics)}"

    def test_term_clause_topic(self):
        recs = sample04_records()
        term = next(r for r in recs if "eleven (11) months" in r.text)
        assert term.topic == "term_renewal"

    def test_signature_witness_kinds(self):
        recs = sample04_records()
        sigs = [r for r in recs if r.kind == "signature"]
        wits = [r for r in recs if r.kind == "witness"]
        assert len(sigs) >= 2 and len(wits) >= 2
        assert all(r.topic == "signature_witness" for r in sigs + wits)

    def test_signatures_not_in_operative_text(self):
        op_text = " ".join(r.text for r in operative(sample04_records()))
        assert "WITNESS 1" not in op_text
        assert "Signature of Landlord" not in op_text

    def test_synthetic_corpus_topics(self):
        # Familiar-wording sanity across the taxonomy (no model needed).
        cases = {
            "The Tenant shall pay monthly rent of Rs. 10,000.": "rent_payment",
            "A refundable security deposit shall be returned within 30 days.": "deposit_refund",
            "Either party may terminate with 30 days notice.": "termination_notice",
            "The Tenant shall vacate and hand over vacant possession on expiry.": "eviction_possession",
            "Late payment attracts a penalty of Rs. 1,000 per month.": "late_penalty",
            "The Landlord shall repair structural defects within 15 days.": "maintenance_repairs",
            "Electricity and water charges shall be paid by the Tenant.": "utilities",
            "The Landlord may enter with 48 hours written notice.": "landlord_access",
            "The Tenant shall not sublet without written consent.": "subletting",
            "The premises shall be used only for residential purposes.": "permitted_use",
            "Disputes shall go to arbitration under the 1996 Act.": "dispute_resolution",
            "This agreement shall remain in force for 11 months.": "term_renewal",
        }
        for text, expected in cases.items():
            got, _ = classify_topic(text)
            assert got == expected, f"{text!r} → {got}, expected {expected}"

    def test_unfamiliar_ownership_wording(self):
        got, _conf = classify_topic(
            "The lessor warrants she holds clear marketable title and is competent to demise the flat.")
        assert got == "ownership_authority"


# ─── 5/6. Risk calibration ─────────────────────────────────────────────────

def _virulence_of(normalised: str, root: str | None = None) -> float:
    from strain.backend.pipeline.analyse import WEIGHTS
    a = assess_asymmetry(normalised)["score"]
    h = assess_harshness_vs_root(normalised, root)["score"]
    o = assess_outcome_factor("no-such-strain", "no-such-family")["value"]
    return WEIGHTS["asymmetry"] * a + WEIGHTS["harshness_delta"] * min(h * 2, 100.0) + WEIGHTS["outcome_factor"] * o


class TestCalibration:
    def test_mutual_notice_not_high(self):
        text = normalise("Either party may terminate this agreement by giving 30 days written notice.")
        assert assess_asymmetry(text)["confidence"] == "high"
        assert _virulence_of(text, text) < 40

    def test_eleven_month_term_not_high_risk(self):
        text = normalise(TERM_11)
        assert _virulence_of(text, text) < 40, _virulence_of(text, text)

    def test_ordinary_rent_low(self):
        text = normalise("The Tenant shall pay a monthly rent of Rs. 18,000 by the 5th of each month.")
        a = assess_asymmetry(text)
        assert a["score"] < 50, a
        assert "penalty" not in " ".join(a["features"])

    def test_refundable_deposit_low(self):
        text = normalise("A refundable deposit shall be returned within 30 days after deducting documented dues.")
        assert assess_asymmetry(text)["score"] < 50

    def test_unilateral_termination_high(self):
        text = normalise("The Landlord may terminate at any time with 7 days notice for any reason the Landlord deems sufficient.")
        assert assess_asymmetry(text)["score"] >= 75

    def test_unrestricted_entry_high(self):
        text = normalise("The Landlord may enter the premises at any time without prior notice.")
        assert assess_asymmetry(text)["score"] >= 75

    def test_penalty_clause_high(self):
        text = normalise("Late rent beyond 7 days attracts a penalty of Rs. 2,000 per month.")
        assert assess_asymmetry(text)["score"] >= 75

    def test_missing_evidence_explicit(self):
        a = assess_asymmetry(normalise("This agreement constitutes the entire understanding."))
        assert a["confidence"] == "low" and a["score"] < 50
        h = assess_harshness_vs_root("some text", None)
        assert h == {"score": 0.0, "confidence": "low", "features": [],
                     "rationale": "No family root available for comparison; no delta asserted."} or h["score"] == 0.0
        o = assess_outcome_factor("no-such-strain", "no-such-family")
        assert o["records"] == 0 and o["verified_sources"] == []

    def test_harshness_needs_real_change(self):
        root = normalise("Either party may terminate with 30 days notice. Breach gets 15 days to cure.")
        same = normalise("Either party may terminate with 30 days notice. Breach gets 15 days to cure.")
        assert assess_harshness_vs_root(same, root)["score"] == 0.0
        harsher = normalise("The Landlord may terminate with 7 days notice. No cure period applies. Breach forfeits the deposit.")
        assert assess_harshness_vs_root(harsher, root)["score"] > 0


# ─── 7/8. Dates ─────────────────────────────────────────────────────────────

class TestKeyDates:
    def test_eleven_month_expiry(self):
        recs = sample04_records()
        entries, anchor = _extract_typed_dates(recs)
        assert anchor and anchor["date"] == date(2026, 4, 10)
        exp = [e for e in entries if e["date_type"] == "expiry"]
        assert exp, entries
        assert exp[0]["date"] == "10 March 2027", exp[0]
        assert exp[0]["basis"] == "calculated"

    def test_execution_not_a_deadline(self):
        recs = sample04_records()
        entries, _ = _extract_typed_dates(recs)
        ex = [e for e in entries if e["date_type"] == "execution"]
        assert len(ex) == 1 and ex[0]["date"] == "10 April 2026"

    def test_missing_anchor_unresolved(self):
        recs = [ClauseRecord("c1", "d", 0, "", TERM_11, normalise(TERM_11))]
        entries, anchor = _extract_typed_dates(recs)
        assert anchor is None
        unres = [e for e in entries if e["basis"] == "unresolved"]
        assert unres and "cannot be determined" in unres[0]["detail"]

    def test_month_end_clamping(self):
        assert _add_duration(date(2026, 1, 31), 1, "months") == date(2026, 2, 28)
        assert _add_duration(date(2024, 1, 31), 1, "months") == date(2024, 2, 29)

    def test_handoff_carries_key_dates(self):
        recs = sample04_records()
        panel = _build_handoff_panel([], recs)
        assert "key_dates" in panel
        labels = [e["label"] for e in panel["key_dates"]]
        assert "Execution date" in labels and "Tenancy expires" in labels


# ─── 9/10. Neutralising selection ────────────────────────────────────────────

def _records(topic: str, text: str, doc: str = "d-src", cid: str = "d-src:c001") -> ClauseRecord:
    return ClauseRecord(cid, doc, 1, "", text, normalise(text), topic=topic, kind="operative")


class TestNeutralising:
    def test_material_terms_compared(self):
        src = extract_clause_features(normalise(TERM_11))
        bad = extract_clause_features(normalise(TERM_3_BAD))
        diffs = _compare_material_terms(src, bad)
        assert any("tenancy length" in d for d in diffs), diffs

    def test_close_variant_passes_material_check(self):
        src = extract_clause_features(normalise(TERM_11))
        near = extract_clause_features(normalise(
            "This agreement shall remain in force for eleven (11) months from the "
            "execution date. Renewal may be made by mutual written consent."))
        assert _compare_material_terms(src, near) == []

    def test_rejects_11_vs_3_month_candidate(self):
        with Session(engine) as session:
            sid = f"test-strain-{uuid.uuid4().hex[:8]}"
            src = _records("term_renewal", TERM_11)
            # Near-identical wording, milder tone, but a 3-month 2011 term.
            cand_text = (
                "This agreement shall remain in force for THREE (3) MONTHS from "
                "15 March 2011. Renewal may be made by mutual written consent.")
            cand = Clause(clause_id=f"{sid}:c1", doc_id=f"{sid}-doc", ordinal=0,
                          heading="", text=cand_text, normalised_text=normalise(cand_text),
                          strain_id=sid, asymmetry_score=15.0)
            session.add(cand)
            session.commit()
            try:
                variant, reason = _select_reference_variant(src, sid, 35.0, session)
                assert variant is None, f"Bad candidate accepted: {variant}"
                assert NO_VARIANT_MESSAGE.split(" — ")[0][:20] in (reason or "")
                assert "tenancy length" in (reason or "").lower()
            finally:
                session.delete(cand)
                session.commit()

    def test_accepts_close_lower_risk_variant(self):
        with Session(engine) as session:
            sid = f"test-strain-{uuid.uuid4().hex[:8]}"
            src = _records("term_renewal", TERM_11)
            cand_text = ("This agreement shall remain in force for eleven (11) months "
                         "from the execution date. Renewal may be made by mutual written consent.")
            cand = Clause(clause_id=f"{sid}:c1", doc_id=f"{sid}-doc", ordinal=0,
                          heading="", text=cand_text, normalised_text=normalise(cand_text),
                          strain_id=sid, asymmetry_score=15.0)
            session.add(cand)
            session.commit()
            try:
                variant, reason = _select_reference_variant(src, sid, 35.0, session)
                assert variant is not None, reason
                assert variant["mode"] == "reference"
                assert variant["material_differences"] == []
            finally:
                session.delete(cand)
                session.commit()


# ─── 11. Litigation disclosure ───────────────────────────────────────────────

class TestLitigationDisclosure:
    def test_records_marked_unverified(self):
        recs = _get_outcomes_for_strain("Security Deposit", "Security Deposit")
        assert recs, "Expected seed outcome records"
        for o in recs:
            assert o["illustrative"] is True
            assert o["verified"] is False
            assert "not a verified judgment" in o["source_disclosure"]

    def test_outcome_evidence_states_basis(self):
        ev = assess_outcome_factor("Security Deposit", "Security Deposit")
        assert ev["basis"] == "illustrative-only"
        assert ev["verified_sources"] == []
        assert "Illustrative" in ev["rationale"] or "illustrative" in ev["rationale"]


# ─── 12/13. Provenance + unclassified ────────────────────────────────────────

class TestProvenance:
    def test_records_carry_source(self):
        for r in sample04_records():
            assert r.clause_id and r.doc_id
            assert r.source, "Missing source locator"

    def test_unclassified_explains_threshold_and_keeps_topic(self):
        cr = _records("term_renewal", TERM_11)
        res = _diagnose_clause(cr, [1.0] + [0.0] * 63, [], Session(engine))
        assert res["status"] == "unclassified"
        assert "0.60" in res["unclassified_message"] or "threshold" in res["unclassified_message"]
        assert res["topic"] == "term_renewal"

    def test_compound_flagged(self):
        assert is_compound("Pay rent monthly. The deposit secures compliance. Initial payment in cash.")
        assert not is_compound("The monthly rent is Rs. 10,000 payable by the 5th.")


# ─── 14/15. Tables, mutation, orientation ────────────────────────────────────

class TestStructure:
    def test_docx_table_extraction(self, tmp_path):
        from docx import Document as DocxDocument
        p = tmp_path / "tabled.docx"
        doc = DocxDocument()
        doc.add_paragraph("RENTAL AGREEMENT")
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "Monthly rent"
        table.cell(0, 1).text = "Rs. 12,000 due by the 5th"
        table.cell(1, 0).text = "Security deposit"
        table.cell(1, 1).text = "Rs. 36,000 refundable"
        doc.add_paragraph("RENTAL AGREEMENT")  # duplicate must not repeat
        doc.save(p)
        text = read_document(p)
        assert "Rs. 12,000" in text and "Rs. 36,000" in text
        assert text.count("RENTAL AGREEMENT") == 1

    def test_mutation_direction_and_types(self):
        assert detect_mutation("pay within 30 days notice", "pay within 7 days notice")[0] == "deadline_change"
        assert detect_mutation("refundable deposit in 30 days", "deposit forfeited in full")[0] == "deposit_change"
        assert detect_mutation("The sky is blue today.", "Quantum flux capacitors hum softly.")[0] == "general_change"

    def test_orientation_dates_and_ties(self):
        from datetime import date as d
        assert orient_edge(d(2020, 1, 1), d(2021, 1, 1), "b", "a") == -1
        assert orient_edge(d(2021, 1, 1), d(2020, 1, 1), "b", "a") == 1
        # Tie → simpler text is parent.
        assert orient_edge(d(2020, 1, 1), d(2020, 1, 1), "plain text", "a. sub\nb. sub\nc. sub") == -1

    def test_shared_headings_do_not_force_kinship(self):
        a = normalise("RENTAL AGREEMENT The landlord owns the building outright.")
        b = normalise("RENTAL AGREEMENT Quantum flux capacitors hum softly near Mars.")
        assert _normalised_edit_distance(a, b) > 0.5
        assert detect_mutation(a, b)[0] == "general_change"

    def test_topic_labels_cover_taxonomy(self):
        for tid in ("term_renewal", "rent_payment", "ownership_authority",
                    "eviction_possession", "signature_witness", "other"):
            assert tid in TOPIC_LABELS


# ─── End-to-end Sample_04 ────────────────────────────────────────────────────

def _stubbed_session():
    return Session(engine)


class TestSample04EndToEnd:
    def test_full_pipeline_response(self, monkeypatch):
        import strain.backend.pipeline.diagnose as diag_mod
        monkeypatch.setattr(diag_mod, "get_embedding_provider",
                            lambda: OfflineEmbeddingProvider())
        stub = OfflineEmbeddingProvider()
        recs = sample04_records()
        term_rec = next(r for r in recs if "eleven (11) months" in r.text)

        sid = f"e2e-strain-{uuid.uuid4().hex[:8]}"
        doc_tmp = f"e2e-doc-{uuid.uuid4().hex[:8]}"
        with Session(engine) as session:
            # Temp strain whose centroid exactly matches the term clause
            # (deterministic stub vectors → distance 0 → classified).
            centroid = stub.encode([term_rec.normalised_text])[0]
            root = Clause(clause_id=f"{doc_tmp}:root", doc_id=doc_tmp, ordinal=0,
                          heading="", text=term_rec.text,
                          normalised_text=term_rec.normalised_text, strain_id=sid,
                          asymmetry_score=35.0)
            session.add(root)
            session.add(Strain(strain_id=sid, family_name="Term Test Family",
                               clause_count=1, root_clause_id=root.clause_id,
                               centroid_json=json.dumps(centroid)))
            # The reported bad candidate: near-identical, 3-month 2011 term.
            bad_text = ("This agreement shall remain in force for THREE (3) MONTHS "
                        "from 15 March 2011. Renewal may be made by mutual written consent.")
            bad = Clause(clause_id=f"{doc_tmp}:bad", doc_id=doc_tmp, ordinal=1,
                         heading="", text=bad_text, normalised_text=normalise(bad_text),
                         strain_id=sid, asymmetry_score=15.0)
            session.add(bad)
            session.commit()
            result: dict = {}
            doc_id_to_clean: str | None = None
            try:
                result = asyncio.run(diagnose(str(FIXTURE_DOCX), "Rental_Agreement_Sample_04.docx", session))
                doc_id_to_clean = result.get("doc_id")
            finally:
                # Cleanup temp rows BEFORE assertions that don't need them.
                for c in session.exec(select(Clause).where(Clause.doc_id == doc_tmp)).all():
                    session.delete(c)
                srow = session.exec(select(Strain).where(Strain.strain_id == sid)).first()
                if srow:
                    session.delete(srow)
                if doc_id_to_clean:
                    for c in session.exec(select(Clause).where(Clause.doc_id == doc_id_to_clean)).all():
                        session.delete(c)
                    drow = session.exec(select(Document).where(Document.doc_id == doc_id_to_clean)).first()
                    if drow:
                        session.delete(drow)
                session.commit()
            assert result, "diagnose() failed before returning a result"

        by_text = {c["text"]: c for c in result["clauses"]}
        term = by_text[term_rec.text]

        # 11-month term: correct topic, not high-risk without evidence.
        assert term["topic"] == "term_renewal", term["topic"]
        assert term["status"] == "classified"
        assert term["virulence_score"] < 55, term["virulence_score"]
        # The 3-month 2011 alternative must be rejected, with the reason shown.
        assert term["neutralising_wording"] is None
        assert term["neutralising_note"] and "No sufficiently similar" in term["neutralising_note"]
        # Expiry calculated or explicitly unresolved.
        key_labels = {e["label"]: e for e in result["handoff_panel"]["key_dates"]}
        assert "Tenancy expires" in key_labels
        exp = key_labels["Tenancy expires"]
        assert exp["date"] == "10 March 2027" or exp["basis"] == "unresolved"
        # No illustrative record posed as verified judgment.
        for c in result["clauses"]:
            for o in c["outcomes"]:
                assert o["verified"] is False
                assert "not a verified judgment" in o["source_disclosure"]
        # Signatures excluded from operative counts.
        assert result["operative_count"] < result["clause_count"]
        assert result["excluded_count"] >= 4
        sig_kinds = {c["kind"] for c in result["clauses"] if c["status"] == "excluded"}
        assert sig_kinds & {"signature", "witness", "preamble"}
        # Lawyer questions match the term topic, not rent.
        term_qs = []
        for item in result["handoff_panel"]["top_risk_clauses"]:
            if item.get("topic") == "term_renewal":
                term_qs = item["lawyer_questions"]
        if term_qs:
            joined = " ".join(term_qs).lower()
            assert "renew" in joined or "tenancy length" in joined or "lock-in" in joined

    def test_provisional_leaf_preserves_lineage(self, monkeypatch):
        import strain.backend.pipeline.diagnose as diag_mod
        monkeypatch.setattr(diag_mod, "get_embedding_provider",
                            lambda: OfflineEmbeddingProvider())
        with Session(engine) as session:
            edges_before = session.exec(select(Edge)).all()
            edge_ids_before = {(e.parent_clause_id, e.child_clause_id) for e in edges_before}
            other_before = {}
            for c in session.exec(select(Clause)).all():
                if c.strain_id and not c.doc_id.startswith("doc-tempcheck"):
                    other_before[c.clause_id] = (c.strain_id, c.virulence_score)
            result = asyncio.run(diagnose(str(FIXTURE_DOCX), "lineage-check.docx", session))
            assert result and "doc_id" in result, "diagnose() failed before returning a result"
            doc_id = result["doc_id"]
            try:
                edges_after = {(e.parent_clause_id, e.child_clause_id)
                               for e in session.exec(select(Edge)).all()}
                assert edges_after == edge_ids_before, "Diagnosis must not rewrite phylogeny edges"
                for c in session.exec(select(Clause)).all():
                    if c.clause_id in other_before:
                        assert (c.strain_id, c.virulence_score) == other_before[c.clause_id], \
                            f"Established clause {c.clause_id} was modified"
                own = session.exec(select(Clause).where(Clause.doc_id == doc_id)).all()
                assert own, "Own clauses must be stored"
            finally:
                for c in session.exec(select(Clause).where(Clause.doc_id == doc_id)).all():
                    session.delete(c)
                drow = session.exec(select(Document).where(Document.doc_id == doc_id)).first()
                if drow:
                    session.delete(drow)
                session.commit()
